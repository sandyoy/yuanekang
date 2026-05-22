#!/usr/bin/env python3
"""提取珠海妇幼数据的核心指标，供白皮书使用"""
import openpyxl
import json
from collections import Counter, defaultdict

wb = openpyxl.load_workbook('data/analysis/珠海妇幼订单数据2021.12-2026.4.xlsx', read_only=True, data_only=True)

output = {}

# ======== 1. 院内+居家汇总 ========
ws = wb['院内+居家']
output['院内居家汇总'] = []
for i, row in enumerate(ws.iter_rows(min_row=3, max_row=7, values_only=True)):
    vals = [str(v)[:30] if v else '-' for v in row[:8]]
    output['院内居家汇总'].append(vals)

# ======== 2. 居家项目占比 ========
ws = wb['居家数据分析']
row3 = list(ws.iter_rows(min_row=3, max_row=3, values_only=True))[0]
row7 = list(ws.iter_rows(min_row=7, max_row=7, values_only=True))[0]
prj_items = []
for j in range(1, 20):
    pname = str(row3[j]).strip() if row3[j] else ''
    pcount = row7[j]
    if pname and pname not in ('年份','订单合计（单）','月均订单','备注','None','') and pcount:
        try:
            cnt = int(float(str(pcount)))
            if cnt > 0:
                prj_items.append((pname, cnt))
        except: pass
total = sum(c for _, c in prj_items)
output['居家项目占比'] = {p: {'单数': c, '占比': f'{c/total*100:.1f}%'} for p, c in sorted(prj_items, key=lambda x: -x[1])}

# ======== 3. 居家明细按月分布 ========
ws = wb['居家明细']
house_monthly = defaultdict(int)
house_projects = Counter()
house_total = 0
for row in ws.iter_rows(min_row=3, values_only=True):
    project = row[3]
    month = row[5]
    if month:
        m_str = str(month).strip()
        if len(m_str) >= 6:
            ym = f"{m_str[:4]}-{m_str[4:6]}"
            house_monthly[ym] += 1
    if project:
        p = str(project).strip()
        if p and p not in ('None', ''):
            house_projects[p] += 1
    house_total += 1

output['居家'] = {
    '总订单': house_total,
    '按月分布': dict(sorted(house_monthly.items())),
    '各项目分布': dict(house_projects.most_common())
}

# ======== 4. 院内明细 ========
ws = wb['院内明细']
inner_monthly = defaultdict(int)
inner_projects = Counter()
inner_total = 0
price_values = []
for row in ws.iter_rows(min_row=2, values_only=True):
    project = row[1]
    price = row[2]
    month = row[6]
    if month:
        m_str = str(month).strip()
        if len(m_str) >= 6:
            ym = f"{m_str[:4]}-{m_str[4:6]}"
            inner_monthly[ym] += 1
    if project:
        p = str(project).strip()
        if p and p not in ('None', ''):
            inner_projects[p] += 1
    if price:
        try:
            pv = float(str(price))
            price_values.append(pv)
        except: pass
    inner_total += 1

# 按年汇总
yearly_inner = defaultdict(int)
for ym, c in inner_monthly.items():
    yearly_inner[ym[:4]] += c

output['院内'] = {
    '总订单': inner_total,
    '按年分布': dict(sorted(yearly_inner.items())),
    '项目TOP15': dict(inner_projects.most_common(15)),
}

if price_values:
    ranges = {'0-99元': 0, '100-299元': 0, '300-499元': 0, '500-999元': 0, '1000元+': 0}
    for v in price_values:
        if v < 100: ranges['0-99元'] += 1
        elif v < 300: ranges['100-299元'] += 1
        elif v < 500: ranges['300-499元'] += 1
        elif v < 1000: ranges['500-999元'] += 1
        else: ranges['1000元+'] += 1
    output['院内']['价格分析'] = {
        '均价': round(sum(price_values)/len(price_values), 0),
        '最高': max(price_values),
        '最低': min(price_values),
        '价格分段': ranges
    }

# ======== 5. 按阶梯计算 ========
all_monthly = defaultdict(int)
for ym, c in house_monthly.items(): all_monthly[ym] += c
for ym, c in inner_monthly.items(): all_monthly[ym] += c

def calc_cost(count):
    if count <= 499: return count * 5.0
    elif count <= 999: return count * 3.5
    elif count <= 1499: return count * 1.5
    else: return count * 1.0

total_examples = 0
total_cost = 0
monthly_detail = []
for ym in sorted(all_monthly.keys()):
    cnt = all_monthly[ym]
    total_examples += cnt
    cost = calc_cost(cnt)
    total_cost += cost
    if cnt <= 499: tier = "0~499例"
    elif cnt <= 999: tier = "500~999例"
    elif cnt <= 1499: tier = "1000~1499例"
    else: tier = "1500例+"
    monthly_detail.append({'月份': ym, '例数': cnt, '档位': tier, '单价': round(cost/cnt, 2), '费用': round(cost, 0)})

output['阶梯结算模拟'] = {
    '累计例数': total_examples,
    '累计费用': round(total_cost, 0),
    '月度明细': monthly_detail
}

# 输出JSON
print(json.dumps(output, ensure_ascii=False, indent=2))
